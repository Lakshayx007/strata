import pandas as pd
import pytest
from openpyxl import load_workbook

from analysis.labeling.build_review import build, excerpt, locate

BODY = " ".join(f"w{i}" for i in range(100)) + " We moved off Redshift because the bill tripled. " + " ".join(f"v{i}" for i in range(100))
DOCS = {1: {"id": 1, "source": "hackernews", "url": "https://news.ycombinator.com/item?id=1", "title": "t", "body": BODY},
        2: {"id": 2, "source": "devto", "url": "https://dev.to/x", "title": None, "body": "Short text.\n\nWe  evaluated BigQuery."}}
TAX = pd.DataFrame([{"code": "cost", "kind": "category", "name": "Cost", "definition": "d", "example_1_doc_id": 1,
                     "example_1_quote": "the bill tripled", "example_2_doc_id": 2, "example_2_quote": "We evaluated BigQuery."},
                    {"code": "none", "kind": "utility", "name": "No reason", "definition": "d", "example_1_doc_id": 2,
                     "example_1_quote": "Short text.", "example_2_doc_id": "", "example_2_quote": ""}])
DRAFTS = pd.DataFrame([
    {"document_id": 1, "taxonomy_code": "cost", "from_vendor": "aws", "to_vendor": "other", "direction": "leave",
     "evidence_span": "We moved off Redshift because the bill tripled."},
    {"document_id": 2, "taxonomy_code": "none", "from_vendor": "none", "to_vendor": "google", "direction": "evaluate",
     "evidence_span": "We evaluated BigQuery."}])


def test_excerpt_is_about_60_words_and_centred():
    e = excerpt(BODY, "We moved off Redshift because the bill tripled.")
    words = e.replace("…", "").split()
    assert len(words) == 60 and "tripled." in e and e.startswith("…") and e.endswith("…")
    assert words.index("We") in range(24, 28)


def test_locate_rejects_paraphrase():
    with pytest.raises(ValueError):
        locate(BODY, "We left Redshift because the bill tripled.")


def test_build_writes_sheets_with_dropdowns(tmp_path):
    out = build(DOCS, TAX, DRAFTS, tmp_path / "r.xlsx")
    wb = load_workbook(out)
    assert {"taxonomy", "labels", "lists"} <= set(wb.sheetnames)
    labels = wb["labels"]
    header = [c.value for c in labels[1]]
    assert header[-5:] == ["my_label", "my_from", "my_to", "my_direction", "notes"]
    assert labels.max_row == 3 and len(labels.data_validations.dataValidation) == 4
    tax = [c.value for c in wb["taxonomy"][1]]
    assert tax[-2:] == ["my_decision", "my_new_name"]


def test_build_refuses_invented_quote(tmp_path):
    bad = TAX.copy()
    bad.loc[0, "example_1_quote"] = "costs went up a lot"
    with pytest.raises(SystemExit):
        build(DOCS, bad, DRAFTS, tmp_path / "r.xlsx")
